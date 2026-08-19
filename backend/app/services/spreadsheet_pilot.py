from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import tempfile
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Callable
from urllib.parse import parse_qsl, urlencode, urljoin, urlparse, urlunparse

import httpx
from openpyxl import load_workbook
from PIL import Image
from sqlalchemy import text
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.config import settings
from app.database import SessionLocal
from app.exceptions import ApiError
from app.models import Brand, Category, ImportItem, Product
from app.services.http_security import build_verified_ssl_context, is_tls_verification_error
from app.services.import_image_ingestion import download_with_retries, validate_downloaded_image
from app.services.imports import is_blocked_source_ip
from app.services.url_security import normalize_hostname, normalized_netloc, parse_ipv4_numeric_literal, resolve_public_hostname, validate_url_authority
from app.utils.slug import make_slug


FANSBUY_PROMOTION_CODE = "59125a2689826be5"
PREFERRED_CODES = ["7792640380", "7792628442", "7792489288", "7789575461", "7789591169", "7792669908", "7789652245"]
CATEGORY_TARGETS = [
    ("01", "sneakers", "Sneakers", 2),
    ("02", "camisetas", "Camisetas", 1),
    ("04", "moletons", "Moletons", 1),
    ("05", "calcas", "Calças", 1),
    ("06", "conjuntos", "Conjuntos", 1),
    ("07", "jaquetas", "Jaquetas", 1),
    ("08", "acessorios", "Acessórios", 1),
]
CATEGORY_MAP = {
    "01": ("sneakers", "Sneakers"),
    "02": ("camisetas", "Camisetas"),
    "03": ("shorts", "Shorts"),
    "04": ("moletons", "Moletons"),
    "05": ("calcas", "Calças"),
    "06": ("conjuntos", "Conjuntos"),
    "07": ("jaquetas", "Jaquetas"),
    "08": ("acessorios", "Acessórios"),
}
BLOCKED_CATEGORY_PREFIXES = {"09", "10", "11", "99"}
BRAND_ALIASES = {
    "lv": "Louis Vuitton",
    "offwhite": "Off-White",
    "off white": "Off-White",
    "casa blanca": "Casablanca",
    "syna": "Syna World",
    "essentials": "Fear of God Essentials",
    "hermes": "Hermès",
    "yeezy": "Adidas",
}
AMBIGUOUS_BRANDS = {"ami"}
KNOWN_PHRASES = [
    "Alexander McQueen",
    "Aimé Leon Dore",
    "Aime Leon Dore",
    "New Balance",
    "Louis Vuitton",
    "The North Face",
    "Stone Island",
    "Palm Angels",
    "Chrome Hearts",
    "Denim Tears",
    "Maison Margiela",
    "Maison Mihara",
    "Golden Goose",
    "Under Armour",
    "Van Cleef & Arpels",
    "Audemars Piguet",
    "Fear of God Essentials",
    "Ralph Lauren",
    "Broken Planet",
    "Gallery Dept",
    "Casa Blanca",
    "Casablanca",
    "Off White",
    "OffWhite",
    "Loro Piana",
    "Moncler",
    "Amiri",
    "Adidas",
    "Nike",
]
GENERIC_NAME_PREFIXES = {"camisas", "camisa", "shirt", "tee", "new"}
PAGE_IMAGE_RE = re.compile(r"https?://[^\"'()<>\\\s]+?\.(?:jpg|jpeg|png|webp)(?:\?[^\"'()<>\\\s]*)?", re.I)
URL_RE = re.compile(r"https?://[^\"']+", re.I)
PROJECT_ROOT = Path(__file__).resolve().parents[3]


@dataclass
class SpreadsheetProduct:
    row: int
    supplier_name: str
    supplier_code: str
    supplier_url: str
    hyperlink_source: str
    image_urls: list[str]
    price_usd_reference: str
    price_brl_reference: str
    category_raw: str
    category_slug: str
    category_name: str
    embedded_images: list[dict] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


@dataclass
class SpreadsheetInspection:
    workbook_path: str
    sheet_name: str
    max_row: int
    max_column: int
    merged_cells: int
    embedded_images: int
    hyperlink_formulas: int
    direct_hyperlinks: int
    image_formulas: int
    candidate_rows: int


def atomic_write_json(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    content = json.dumps(payload, ensure_ascii=False, indent=2)
    json.loads(content)
    with tempfile.NamedTemporaryFile("w", encoding="utf-8", dir=path.parent, prefix=f".{path.name}.", suffix=".tmp", delete=False) as handle:
        tmp_path = Path(handle.name)
        handle.write(content)
        handle.flush()
    tmp_path.replace(path)


def workbook_sha(path: Path) -> str:
    hasher = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            hasher.update(chunk)
    return hasher.hexdigest()


def clean_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def supplier_code(value) -> str:
    text = clean_text(value)
    return re.sub(r"\.0$", "", text)


def extract_urls_from_formula(value) -> list[str]:
    if not isinstance(value, str) or not value.startswith("="):
        return []
    return URL_RE.findall(value)


def image_urls_from_cell(value) -> list[str]:
    if not isinstance(value, str) or not value.startswith("="):
        return []
    if "IMAGE" not in value.upper():
        return []
    return [url for url in extract_urls_from_formula(value) if urlparse(url).scheme == "https"]


def embedded_images_by_row(ws) -> dict[int, list[dict]]:
    result: dict[int, list[dict]] = {}
    for image in getattr(ws, "_images", []):
        marker = getattr(getattr(image, "anchor", None), "_from", None)
        if marker is None:
            continue
        try:
            content = image._data()
        except Exception:
            continue
        row = marker.row + 1
        result.setdefault(row, []).append(
            {
                "content": content,
                "format": str(getattr(image, "format", "") or "").lower(),
                "anchor": {"row": row, "column": marker.col + 1},
                "width": getattr(image, "width", None),
                "height": getattr(image, "height", None),
            }
        )
    return result


def hyperlink_url(cell, code: str) -> tuple[str, str]:
    if cell.hyperlink and cell.hyperlink.target:
        return str(cell.hyperlink.target), "direct_hyperlink"
    value = cell.value
    if isinstance(value, str) and "HYPERLINK" in value.upper():
        urls = extract_urls_from_formula(value)
        if urls and '"&$F' not in value and "'&$F" not in value and "&$F" not in value:
            return urls[0], "hyperlink_formula_literal"
        if code and re.search(r'HYPERLINK\("https://fansbuy\.com/item-micro-"&\$F\d+&"\.html\?promotionCode=59125a2689826be5"', value):
            return fansbuy_url_from_code(code), "fansbuy_formula_derived"
    if code:
        return fansbuy_url_from_code(code), "fansbuy_code_derived"
    return "", "missing"


def fansbuy_url_from_code(code: str) -> str:
    return f"https://fansbuy.com/item-micro-{code}.html?promotionCode={FANSBUY_PROMOTION_CODE}"


def category_from_raw(raw: str) -> tuple[str, str]:
    prefix = (raw or "").strip()[:2]
    return CATEGORY_MAP.get(prefix, ("", ""))


def normalize_url(url: str) -> str:
    parsed = urlparse(str(url).strip())
    query = [(key, value) for key, value in parse_qsl(parsed.query, keep_blank_values=True) if key.lower() not in {"utm_source", "utm_medium", "utm_campaign"}]
    return urlunparse(parsed._replace(scheme=parsed.scheme.lower(), netloc=(parsed.netloc or "").lower().rstrip("."), fragment="", query=urlencode(sorted(query))))


def resolve_hostname(hostname: str) -> list[str]:
    hostname = normalize_hostname(hostname)
    numeric = parse_ipv4_numeric_literal(hostname)
    if numeric:
        return [numeric]
    try:
        return resolve_public_hostname(hostname)
    except ApiError as exc:
        raise ValueError(f"blocked_host:{hostname}") from exc


def validate_public_http_url(url: str) -> str:
    parsed = urlparse(str(url).strip())
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        raise ValueError("unsupported_scheme")
    try:
        validate_url_authority(url, code="IMPORT_INVALID_URL", message="Informe uma URL HTTP ou HTTPS valida.")
    except ApiError as exc:
        blocker = exc.details.get("blocker") if isinstance(exc.details, dict) else "invalid_url"
        raise ValueError(blocker) from exc
    host = normalize_hostname(parsed.hostname)
    for address in resolve_hostname(host):
        if is_blocked_source_ip(address):
            raise ValueError("blocked_ip")
    return parsed._replace(netloc=normalized_netloc(host, parsed.port), fragment="").geturl()


def load_spreadsheet_products(xlsx_path: Path) -> tuple[SpreadsheetInspection, list[SpreadsheetProduct]]:
    formula_wb = load_workbook(xlsx_path, data_only=False, read_only=False)
    value_wb = load_workbook(xlsx_path, data_only=True, read_only=False)
    ws = formula_wb.active
    value_ws = value_wb[ws.title]
    embedded_by_row = embedded_images_by_row(ws)
    products: list[SpreadsheetProduct] = []
    hyperlink_formulas = direct_hyperlinks = image_formulas = 0
    for row in range(1, ws.max_row + 1):
        name = clean_text(ws.cell(row, 1).value)
        if not name or name.upper() in {"PRODUTO"}:
            continue
        code = supplier_code(ws.cell(row, 6).value)
        category_raw = clean_text(value_ws.cell(row, 7).value or ws.cell(row, 7).value)
        category_prefix = category_raw[:2]
        if category_prefix in BLOCKED_CATEGORY_PREFIXES or category_raw == "99 REMOVER":
            continue
        slug, category_name = category_from_raw(category_raw)
        if not slug:
            continue
        link_cell = ws.cell(row, 3)
        if link_cell.hyperlink:
            direct_hyperlinks += 1
        if isinstance(link_cell.value, str) and "HYPERLINK" in link_cell.value.upper():
            hyperlink_formulas += 1
        image_urls = image_urls_from_cell(ws.cell(row, 2).value)
        if image_urls:
            image_formulas += 1
        supplier_url, source = hyperlink_url(link_cell, code)
        if not supplier_url:
            continue
        products.append(
            SpreadsheetProduct(
                row=row,
                supplier_name=name,
                supplier_code=code,
                supplier_url=supplier_url,
                hyperlink_source=source,
                image_urls=image_urls,
                price_usd_reference=clean_text(ws.cell(row, 4).value),
                price_brl_reference=clean_text(ws.cell(row, 5).value),
                category_raw=category_raw,
                category_slug=slug,
                category_name=category_name,
                embedded_images=embedded_by_row.get(row, []),
            )
        )
    inspection = SpreadsheetInspection(
        workbook_path=str(xlsx_path),
        sheet_name=ws.title,
        max_row=ws.max_row,
        max_column=ws.max_column,
        merged_cells=len(ws.merged_cells.ranges),
        embedded_images=len(getattr(ws, "_images", [])),
        hyperlink_formulas=hyperlink_formulas,
        direct_hyperlinks=direct_hyperlinks,
        image_formulas=image_formulas,
        candidate_rows=len(products),
    )
    return inspection, products


def select_pilot_products(products: list[SpreadsheetProduct], *, limit: int, rows: set[int] | None = None, codes: set[str] | None = None) -> list[SpreadsheetProduct]:
    if limit > 10:
        raise ValueError("O piloto nao pode processar mais de 10 produtos.")
    selected: list[SpreadsheetProduct] = []
    seen: set[str] = set()

    def add(product: SpreadsheetProduct) -> None:
        key = product.supplier_code or normalize_url(product.supplier_url)
        if key not in seen and len(selected) < limit:
            selected.append(product)
            seen.add(key)

    if rows or codes:
        for product in products:
            if (rows and product.row in rows) or (codes and product.supplier_code in codes):
                add(product)
        return selected[:limit]

    by_code = {product.supplier_code: product for product in products if product.supplier_code}
    for code in PREFERRED_CODES:
        if code in by_code:
            add(by_code[code])
    for prefix, slug, _name, target in CATEGORY_TARGETS:
        while sum(1 for item in selected if item.category_slug == slug) < target and len(selected) < limit:
            candidate = next((item for item in products if item.category_raw.startswith(prefix) and (item.supplier_code or normalize_url(item.supplier_url)) not in seen), None)
            if not candidate:
                break
            add(candidate)
    for product in products:
        if len(selected) >= limit:
            break
        add(product)
    return selected


def canonical_brand_name(name: str) -> str:
    first_line = re.sub(r"\[[^\]]+\]", "", (name or "").splitlines()[0]).strip()
    words_for_cleanup = first_line.split()
    while words_for_cleanup and (words_for_cleanup[0].casefold() in GENERIC_NAME_PREFIXES or words_for_cleanup[0].isdigit()):
        words_for_cleanup.pop(0)
    first_line = " ".join(words_for_cleanup) or first_line
    folded = re.sub(r"\s+", " ", first_line).casefold()
    for alias, canonical in BRAND_ALIASES.items():
        if folded == alias or folded.startswith(f"{alias} ") or folded.startswith(f"{alias} x "):
            return canonical
    for phrase in sorted(KNOWN_PHRASES, key=len, reverse=True):
        if folded.startswith(phrase.casefold()):
            return BRAND_ALIASES.get(phrase.casefold(), phrase)
    words = first_line.split()
    if not words:
        return ""
    first = words[0].strip("×x/&")
    if first.casefold() == "air" and len(words) > 1:
        return "Jordan" if words[1].casefold() == "jordan" else "Nike"
    if first.casefold() == "jordan":
        return "Jordan"
    return BRAND_ALIASES.get(first.casefold(), first)


def is_ambiguous_brand(name: str, product_name: str) -> bool:
    folded = name.casefold()
    first_line = product_name.splitlines()[0].casefold()
    if folded in AMBIGUOUS_BRANDS:
        return True
    return " x " in first_line and name not in {"Nike", "Adidas", "Jordan", "Louis Vuitton", "Syna World"}


def static_brand_lookup() -> dict[str, object]:
    path = PROJECT_ROOT / "frontend" / "data" / "brands.json"
    if not path.is_file():
        return {}
    data = json.loads(path.read_text(encoding="utf-8-sig"))
    return {str(item.get("name", "")).casefold(): type("StaticBrand", (), {"id": None, "name": item.get("name")})() for item in data.get("brands", [])}


def static_category_lookup() -> dict[str, object]:
    path = PROJECT_ROOT / "frontend" / "data" / "categories.json"
    if not path.is_file():
        return {}
    data = json.loads(path.read_text(encoding="utf-8-sig"))
    result = {}
    for key in ("clothing", "all"):
        for item in data.get(key, []):
            slug = item.get("slug")
            if slug and slug not in result:
                result[slug] = type("StaticCategory", (), {"id": None, "slug": slug, "name": item.get("name")})()
    return result


def brand_lookup(db: Session | None) -> dict[str, object]:
    if db is None:
        return static_brand_lookup()
    return {brand.name.casefold(): brand for brand in db.query(Brand).all()}


def category_lookup(db: Session | None) -> dict[str, object]:
    if db is None:
        return static_category_lookup()
    return {category.slug: category for category in db.query(Category).all()}


def duplicate_check(db: Session | None, product: SpreadsheetProduct, selected: list[SpreadsheetProduct]) -> dict:
    code = product.supplier_code
    normalized_url = normalize_url(product.supplier_url)
    pilot_rows = [item.row for item in selected if item is not product and ((code and item.supplier_code == code) or normalize_url(item.supplier_url) == normalized_url)]
    if db is None:
        status = "possible_duplicate" if pilot_rows else "manual_review"
        return {"supplier_code": code, "normalized_url": normalized_url, "pilot_duplicate_rows": pilot_rows, "import_item_ids": [], "product_ids_by_sku": [], "source_url_item_ids": [], "database_checked": False, "status": status, "is_duplicate": bool(pilot_rows)}
    import_items = []
    if code:
        import_items = [item.id for item in db.query(ImportItem).filter(ImportItem.external_id == code).all()]
    source_url_items = [item.id for item in db.query(ImportItem).filter(ImportItem.normalized_source_url == normalized_url).all()]
    product_skus = [item.id for item in db.query(Product).filter(Product.sku == code).all()] if code else []
    if pilot_rows:
        status = "possible_duplicate"
    elif import_items:
        status = "existing_supplier_code"
    elif source_url_items:
        status = "existing_source_url"
    elif product_skus:
        status = "exact_duplicate"
    else:
        status = "new"
    return {"supplier_code": code, "normalized_url": normalized_url, "pilot_duplicate_rows": pilot_rows, "import_item_ids": import_items, "product_ids_by_sku": product_skus, "source_url_item_ids": source_url_items, "database_checked": True, "status": status, "is_duplicate": bool(pilot_rows or import_items or source_url_items or product_skus)}


def extract_page_data(html: str, final_url: str) -> dict:
    title_match = re.search(r"<title[^>]*>(.*?)</title>", html, re.I | re.S)
    title = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", "", title_match.group(1))).strip() if title_match else ""
    image_urls = []
    for url in PAGE_IMAGE_RE.findall(html):
        absolute = urljoin(final_url, url)
        if absolute not in image_urls:
            image_urls.append(absolute)
    lower = html.casefold()
    blocked = any(term in lower for term in ["login", "sign in", "captcha", "access denied", "forbidden"])
    return {"title": title, "image_urls": image_urls[:8], "login_or_blocked": blocked}


def fetch_page(url: str, client_factory: Callable[[], httpx.Client] | None = None) -> dict:
    try:
        current_url = validate_public_http_url(url)
    except ValueError as exc:
        return {"status": "blocked", "error": str(exc), "final_url": url, "domain": urlparse(url).hostname or ""}
    redirects = 0
    try:
        factory = client_factory or (lambda: httpx.Client(timeout=10, follow_redirects=False, headers={"User-Agent": settings.import_user_agent}, verify=build_verified_ssl_context()))
        with factory() as client:
            while True:
                with client.stream("GET", current_url) as response:
                    if response.status_code in {301, 302, 303, 307, 308}:
                        location = response.headers.get("location")
                        if not location:
                            return {"status": "failed", "status_code": response.status_code, "error": "redirect_without_location", "final_url": current_url, "domain": urlparse(current_url).hostname or ""}
                        redirects += 1
                        if redirects > 3:
                            return {"status": "failed", "status_code": response.status_code, "error": "too_many_redirects", "final_url": current_url, "domain": urlparse(current_url).hostname or ""}
                        current_url = validate_public_http_url(urljoin(current_url, location))
                        continue
                    content = bytearray()
                    for chunk in response.iter_bytes(64 * 1024):
                        if not chunk:
                            continue
                        content.extend(chunk)
                        if len(content) > settings.import_max_response_bytes:
                            return {"status": "failed", "status_code": response.status_code, "error": "response_too_large", "final_url": current_url, "domain": urlparse(current_url).hostname or ""}
                    body = bytes(content).decode(response.encoding or "utf-8", errors="replace")
                    content_type = response.headers.get("content-type", "").split(";")[0].lower()
                    result = {"status": "ok" if response.status_code < 400 else "failed", "status_code": response.status_code, "final_url": str(response.url), "domain": urlparse(str(response.url)).hostname or "", "content_type": content_type}
                    if content_type.startswith("text/html"):
                        result.update(extract_page_data(body, str(response.url)))
                    return result
    except httpx.HTTPError as exc:
        error = "tls_verification_failed" if is_tls_verification_error(exc) else "http_error"
        return {"status": "failed", "error": error, "detail": str(exc)[:200], "final_url": current_url, "domain": urlparse(current_url).hostname or ""}
    except (ValueError, OSError) as exc:
        return {"status": "blocked", "error": str(exc), "final_url": current_url, "domain": urlparse(current_url).hostname or ""}


def download_pilot_image(url: str, images_dir: Path, *, supplier_code: str, position: int, referer: str | None = None, client_factory: Callable[[], httpx.Client] | None = None) -> dict:
    images_dir.mkdir(parents=True, exist_ok=True)
    try:
        timeout = httpx.Timeout(connect=settings.import_image_connect_timeout, read=settings.import_image_read_timeout, write=settings.import_image_read_timeout, pool=settings.import_image_connect_timeout)
        factory = client_factory or (lambda: httpx.Client(timeout=timeout, follow_redirects=False, verify=build_verified_ssl_context()))
        with factory() as client:
            temp_path, content_hash, header_mime, size_bytes = download_with_retries(client, url, images_dir, referer=referer)
        mime, width, height, warnings = validate_downloaded_image(temp_path, header_mime)
        extension = {"image/jpeg": ".jpg", "image/png": ".png", "image/webp": ".webp"}[mime]
        filename = f"{supplier_code or 'no-code'}-{position:02d}-{content_hash[:16]}{extension}"
        final_path = images_dir / filename
        if not final_path.exists():
            shutil.move(str(temp_path), final_path)
            created = True
        else:
            temp_path.unlink(missing_ok=True)
            created = False
        return {"source_type": "url", "source_url": url, "source_domain": urlparse(url).hostname or "", "position": position, "filename": filename, "path": str(final_path), "mime": mime, "size_bytes": size_bytes, "sha256": content_hash, "width": width, "height": height, "status": "stored", "created": created, "classification": "low_resolution" if "small_image" in warnings else "supplier_usable", "warnings": warnings}
    except Exception as exc:
        return {"source_type": "url", "source_url": url, "source_domain": urlparse(url).hostname or "", "position": position, "status": "failed", "error": str(exc)[:220], "classification": "missing"}


def store_embedded_pilot_image(image: dict, images_dir: Path, *, supplier_code: str, position: int) -> dict:
    images_dir.mkdir(parents=True, exist_ok=True)
    content = image.get("content") or b""
    content_hash = hashlib.sha256(content).hexdigest()
    header_mime = {"png": "image/png", "jpeg": "image/jpeg", "jpg": "image/jpeg", "webp": "image/webp"}.get(image.get("format"))
    temp_path = None
    try:
        with tempfile.NamedTemporaryFile("wb", dir=images_dir, prefix=".embedded-", suffix=".tmp", delete=False) as handle:
            temp_path = Path(handle.name)
            handle.write(content)
        mime, width, height, warnings = validate_downloaded_image(temp_path, header_mime)
        extension = {"image/jpeg": ".jpg", "image/png": ".png", "image/webp": ".webp"}[mime]
        filename = f"{supplier_code or 'no-code'}-{position:02d}-{content_hash[:16]}{extension}"
        final_path = images_dir / filename
        if not final_path.exists():
            shutil.move(str(temp_path), final_path)
            created = True
        else:
            temp_path.unlink(missing_ok=True)
            created = False
        return {"source_type": "embedded_xlsx", "source_url": None, "source_domain": "", "anchor": image.get("anchor"), "position": position, "filename": filename, "path": str(final_path), "mime": mime, "size_bytes": len(content), "sha256": content_hash, "width": width, "height": height, "status": "stored", "created": created, "classification": "low_resolution" if "small_image" in warnings else "supplier_usable", "warnings": warnings}
    except Exception as exc:
        if temp_path is not None:
            temp_path.unlink(missing_ok=True)
        return {"source_type": "embedded_xlsx", "source_url": None, "source_domain": "", "anchor": image.get("anchor"), "position": position, "status": "failed", "error": str(exc)[:220], "classification": "missing"}


def suggested_product_name(product: SpreadsheetProduct, brand_name: str) -> str:
    first_line = re.sub(r"\[[^\]]+\]", "", product.supplier_name.splitlines()[0]).strip()
    folded = first_line.casefold()
    if brand_name and (folded.startswith(brand_name.casefold()) or (brand_name == "Jordan" and "jordan" in folded)):
        return first_line
    return f"{brand_name} {first_line}".strip()


def build_proposal(db: Session | None, product: SpreadsheetProduct, selected: list[SpreadsheetProduct], run_dirs: dict[str, Path], *, dry_run: bool, fetch: bool = True, client_factory=None, db_warning: str | None = None) -> dict:
    warnings: list[str] = list(product.warnings)
    if db_warning:
        warnings.append(db_warning)
    category_by_slug = category_lookup(db)
    brand_by_name = brand_lookup(db)
    brand_name = canonical_brand_name(product.supplier_name)
    brand = brand_by_name.get(brand_name.casefold()) if brand_name else None
    ambiguous = is_ambiguous_brand(brand_name, product.supplier_name) if brand_name else True
    if brand:
        brand_status = "existing"
    elif ambiguous:
        brand_status = "manual_review"
        warnings.append("brand_requires_manual_review")
    else:
        brand_status = "would_create" if dry_run else "created"
    category = category_by_slug.get(product.category_slug)
    if not category:
        warnings.append("category_not_found_in_db")
    dup = duplicate_check(db, product, selected)
    if dup["is_duplicate"]:
        warnings.append("possible_duplicate")
    page = fetch_page(product.supplier_url, client_factory=client_factory) if fetch else {"status": "skipped", "domain": urlparse(product.supplier_url).hostname or ""}
    if page.get("status") != "ok":
        warnings.append("supplier_page_not_accessible")
    if page.get("login_or_blocked"):
        warnings.append("supplier_page_login_or_blocked")
    image_sources = list(dict.fromkeys(product.image_urls + page.get("image_urls", [])))[:3]
    downloaded = [download_pilot_image(url, run_dirs["images"], supplier_code=product.supplier_code, position=index + 1, referer=product.supplier_url, client_factory=client_factory) for index, url in enumerate(image_sources)]
    remaining_slots = max(0, 3 - len(downloaded))
    downloaded.extend(
        store_embedded_pilot_image(image, run_dirs["images"], supplier_code=product.supplier_code, position=len(downloaded) + index + 1)
        for index, image in enumerate(product.embedded_images[:remaining_slots])
    )
    if not downloaded:
        warnings.append("no_images_found")
    if any(item.get("status") != "stored" for item in downloaded):
        warnings.append("image_download_incomplete")
    return {
        "spreadsheet_row": product.row,
        "supplier_code": product.supplier_code,
        "supplier_url": product.supplier_url,
        "hyperlink_source": product.hyperlink_source,
        "supplier_name": product.supplier_name,
        "suggested_name": suggested_product_name(product, brand_name),
        "detected_brand": brand_name,
        "brand_id": getattr(brand, "id", None) if brand else None,
        "brand_status": brand_status,
        "category_id": getattr(category, "id", None) if category else None,
        "category_slug": product.category_slug,
        "category_name": product.category_name,
        "spreadsheet_category": product.category_raw,
        "price_usd_reference": product.price_usd_reference,
        "price_brl_reference": product.price_brl_reference,
        "commercial_price": None,
        "supplier_page": page,
        "supplier_images": downloaded,
        "alternative_image_candidates": [],
        "duplicate_check": dup,
        "warnings": list(dict.fromkeys(warnings)),
        "review_required": True,
    }


def unique_run_dir(output_root: Path, base_run_id: str) -> tuple[str, Path]:
    run_dir = output_root / base_run_id
    if not run_dir.exists():
        return base_run_id, run_dir
    suffix = datetime.now(timezone.utc).strftime("%Y%m%d%H%M%S")
    candidate = output_root / f"{base_run_id}-{suffix}"
    counter = 2
    while candidate.exists():
        candidate = output_root / f"{base_run_id}-{suffix}-{counter}"
        counter += 1
    return candidate.name, candidate


def run_spreadsheet_pilot(*, xlsx: Path, limit: int, dry_run: bool, rows: set[int] | None = None, supplier_codes: set[str] | None = None, verbose: bool = False, output_root: Path = PROJECT_ROOT / "imports" / "pilots", client_factory=None) -> dict:
    if limit > 10:
        raise ValueError("O piloto nao pode processar mais de 10 produtos.")
    digest = workbook_sha(xlsx)
    base_run_id = f"pilot-{digest[:8]}-limit{limit}"
    run_id, run_dir = unique_run_dir(output_root, base_run_id)
    run_dirs = {"root": run_dir, "products": run_dir / "products", "images": run_dir / "images", "logs": run_dir / "logs"}
    for directory in run_dirs.values():
        directory.mkdir(parents=True, exist_ok=True)
    inspection, products = load_spreadsheet_products(xlsx)
    selected = select_pilot_products(products, limit=limit, rows=rows, codes=supplier_codes)
    db_warning = None
    db: Session | None = None
    session_cm = None
    try:
        session_cm = SessionLocal()
        session_cm.execute(text("select 1"))
        db = session_cm
    except SQLAlchemyError:
        if session_cm is not None:
            session_cm.close()
        db = None
        db_warning = "database_unavailable_duplicate_check_skipped"
    try:
        proposals = [build_proposal(db, product, selected, run_dirs, dry_run=dry_run, client_factory=client_factory, db_warning=db_warning) for product in selected]
        for proposal in proposals:
            atomic_write_json(run_dirs["products"] / f"{proposal['spreadsheet_row']}-{proposal['supplier_code'] or 'no-code'}.json", proposal)
        manifest = {
            "run_id": run_id,
            "base_run_id": base_run_id,
            "dry_run": dry_run,
            "database_connected": db is not None,
            "created_at": datetime.now(timezone.utc).isoformat(),
            "xlsx": str(xlsx),
            "workbook_sha256": digest,
            "inspection": inspection.__dict__,
            "limit": limit,
            "selected_rows": [item.row for item in selected],
            "selected_supplier_codes": [item.supplier_code for item in selected],
            "products_count": len(proposals),
            "links_accessible": sum(1 for item in proposals if item["supplier_page"].get("status") == "ok"),
            "links_login_or_blocked": sum(1 for item in proposals if item["supplier_page"].get("login_or_blocked")),
            "images_found": sum(len(item["supplier_images"]) for item in proposals),
            "images_downloaded": sum(1 for item in proposals for image in item["supplier_images"] if image.get("status") == "stored"),
            "brand_status_counts": dict(sorted({status: sum(1 for item in proposals if item["brand_status"] == status) for status in {item["brand_status"] for item in proposals}}.items())),
            "warnings": sorted({warning for item in proposals for warning in item["warnings"]}),
            "proposals": [f"products/{proposal['spreadsheet_row']}-{proposal['supplier_code'] or 'no-code'}.json" for proposal in proposals],
        }
        atomic_write_json(run_dirs["root"] / "manifest.json", manifest)
    finally:
        if db is not None:
            db.close()
    if verbose:
        print(json.dumps(manifest, ensure_ascii=False, indent=2))
    return manifest


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Piloto controlado de importacao da planilha DripZone.")
    parser.add_argument("--xlsx", required=True)
    parser.add_argument("--limit", type=int, default=10)
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--row", action="append", type=int, default=[])
    parser.add_argument("--supplier-code", action="append", default=[])
    parser.add_argument("--verbose", action="store_true")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    manifest = run_spreadsheet_pilot(
        xlsx=Path(args.xlsx),
        limit=args.limit,
        dry_run=args.dry_run,
        rows=set(args.row) if args.row else None,
        supplier_codes=set(args.supplier_code) if args.supplier_code else None,
        verbose=args.verbose,
    )
    if not args.verbose:
        print(json.dumps({"run_id": manifest["run_id"], "products_count": manifest["products_count"], "manifest": str(Path("imports/pilots") / manifest["run_id"] / "manifest.json")}, ensure_ascii=False))
    return 0
