from __future__ import annotations

import json
from pathlib import Path

import httpx
import pytest
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlsxImage
from PIL import Image
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.database import Base
from app.models import Brand, Category
from app.services import spreadsheet_pilot as pilot


def image_bytes(path: Path) -> bytes:
    Image.new("RGB", (20, 20), "white").save(path, format="PNG")
    return path.read_bytes()


def make_workbook(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Mais Vendidos"
    ws.append(["PRODUTO", "IMAGEM", "LINK", "PREÇO (USD)", "PREÇO (BRL)", "CODE", "CATEGORY"])
    rows = [
        ("Air Jordan 4\n(Várias cores)", '=IMAGE("https://img.example/jordan.png")', '=IF($F2="","",HYPERLINK("https://fansbuy.com/item-micro-"&$F2&".html?promotionCode=59125a2689826be5","LINK"))', "$80", "R$430", 7792640380, "01 👟 TÊNIS"),
        ("Adidas Campus [PK BATCH]\n(Várias cores)", "", '=HYPERLINK("https://fansbuy.com/item-micro-7792628442.html?promotionCode=59125a2689826be5","LINK")', "$37", "R$196", 7792628442, "01 👟 TÊNIS"),
        ("Camisas Ami\n(Várias cores)", "", '=IF($F4="","",HYPERLINK("https://fansbuy.com/item-micro-"&$F4&".html?promotionCode=59125a2689826be5","LINK"))', "$16", "R$84", 7789575461, "02 👕 CAMISETAS E POLOS"),
        ("Adwysd Hoodie\n22 opções disponíveis", "", '=IF($F5="","",HYPERLINK("https://fansbuy.com/item-micro-"&$F5&".html?promotionCode=59125a2689826be5","LINK"))', "$26", "R$136", 7792669908, "04 🥶 MOLETONS E HOODIES"),
        ("Nike Pants", "", '=IF($F6="","",HYPERLINK("https://fansbuy.com/item-micro-"&$F6&".html?promotionCode=59125a2689826be5","LINK"))', "$20", "R$100", 1111111111, "05 👖 CALÇAS E JEANS"),
        ("Trapstar Set", "", '=IF($F7="","",HYPERLINK("https://fansbuy.com/item-micro-"&$F7&".html?promotionCode=59125a2689826be5","LINK"))', "$22", "R$112", 2222222222, "06 🎽 CONJUNTOS"),
        ("The North Face Jacket", "", '=IF($F8="","",HYPERLINK("https://fansbuy.com/item-micro-"&$F8&".html?promotionCode=59125a2689826be5","LINK"))', "$30", "R$150", 3333333333, "07 🧥 JAQUETAS E COLETES"),
        ("Adidas Bag", "", '=IF($F9="","",HYPERLINK("https://fansbuy.com/item-micro-"&$F9&".html?promotionCode=59125a2689826be5","LINK"))', "$10", "R$52", 7789652245, "08 👜 ACESSÓRIOS"),
        ("Armani Perfume", "", '=IF($F10="","",HYPERLINK("https://fansbuy.com/item-micro-"&$F10&".html?promotionCode=59125a2689826be5","LINK"))', "$10", "R$52", 9999999999, "10 🌹 PERFUMES"),
    ]
    for row in rows:
        ws.append(row)
    ws["C11"] = "LINK"
    ws["C11"].hyperlink = "https://fansbuy.com/item-micro-4444444444.html?promotionCode=59125a2689826be5"
    ws["A11"] = "OffWhite Tee"
    ws["D11"] = "$12"
    ws["E11"] = "R$60"
    ws["F11"] = 4444444444
    ws["G11"] = "02 👕 CAMISETAS E POLOS"
    wb.save(path)
    return path


@pytest.fixture()
def pilot_db(monkeypatch, tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path / 'pilot.db'}", connect_args={"check_same_thread": False})
    Session = sessionmaker(bind=engine, autoflush=False, autocommit=False, expire_on_commit=False)
    Base.metadata.create_all(bind=engine)
    with Session() as db:
        for name, slug in [
            ("Adidas", "adidas"),
            ("Jordan", "jordan"),
            ("Nike", "nike"),
            ("Off-White", "off-white"),
            ("Fear of God Essentials", "fear-of-god-essentials"),
            ("Syna World", "syna-world"),
            ("Trapstar", "trapstar"),
            ("The North Face", "the-north-face"),
        ]:
            db.add(Brand(name=name, slug=slug, is_active=True))
        for name, slug in [
            ("Sneakers", "sneakers"),
            ("Camisetas", "camisetas"),
            ("Moletons", "moletons"),
            ("Calças", "calcas"),
            ("Conjuntos", "conjuntos"),
            ("Jaquetas", "jaquetas"),
            ("Acessórios", "acessorios"),
        ]:
            db.add(Category(name=name, slug=slug, is_active=True))
        db.commit()
    monkeypatch.setattr(pilot, "SessionLocal", Session)
    return Session


def fake_client_factory(tmp_path):
    body_image = image_bytes(tmp_path / "fixture.png")

    def handler(request: httpx.Request) -> httpx.Response:
        if request.url.host == "img.example":
            return httpx.Response(200, content=body_image, headers={"content-type": "image/png"}, request=request)
        html = '<html><head><title>Fansbuy Item</title></head><body><img src="https://img.example/page.png"></body></html>'
        return httpx.Response(200, text=html, headers={"content-type": "text/html"}, request=request)

    return lambda: httpx.Client(transport=httpx.MockTransport(handler), follow_redirects=False)


def test_loads_xlsx_hyperlink_formula_direct_hyperlink_image_and_category(tmp_path, monkeypatch):
    monkeypatch.setattr(pilot, "resolve_hostname", lambda host: ["93.184.216.34"])
    path = make_workbook(tmp_path / "pilot.xlsx")
    inspection, products = pilot.load_spreadsheet_products(path)
    by_code = {item.supplier_code: item for item in products}
    assert inspection.sheet_name == "Mais Vendidos"
    assert inspection.hyperlink_formulas >= 8
    assert inspection.image_formulas == 1
    assert by_code["7792640380"].supplier_url == "https://fansbuy.com/item-micro-7792640380.html?promotionCode=59125a2689826be5"
    assert by_code["4444444444"].hyperlink_source == "direct_hyperlink"
    assert by_code["7792640380"].image_urls == ["https://img.example/jordan.png"]
    assert by_code["1111111111"].category_slug == "calcas"
    assert "9999999999" not in by_code


def test_brand_alias_ambiguity_and_limit_validation():
    assert pilot.canonical_brand_name("LV Bracelet") == "Louis Vuitton"
    assert pilot.canonical_brand_name("Essentials Hoodie") == "Fear of God Essentials"
    assert pilot.canonical_brand_name("OffWhite Tee") == "Off-White"
    assert pilot.is_ambiguous_brand("Ami", "Ami Shirt")
    with pytest.raises(ValueError):
        pilot.select_pilot_products([], limit=11)


def test_url_security_blocks_unsafe(monkeypatch):
    monkeypatch.setattr(pilot, "resolve_hostname", lambda host: ["127.0.0.1"])
    with pytest.raises(ValueError):
        pilot.validate_public_http_url("https://localhost/item")
    with pytest.raises(ValueError):
        pilot.validate_public_http_url("javascript:alert(1)")


def test_dry_run_generates_idempotent_proposals_without_commercial_price(tmp_path, monkeypatch, pilot_db):
    monkeypatch.setattr(pilot, "resolve_hostname", lambda host: ["93.184.216.34"])
    monkeypatch.setattr("app.services.import_image_ingestion.resolve_hostname", lambda host: ["93.184.216.34"])
    path = make_workbook(tmp_path / "pilot.xlsx")
    output = tmp_path / "pilots"
    manifest = pilot.run_spreadsheet_pilot(xlsx=path, limit=5, dry_run=True, output_root=output, client_factory=fake_client_factory(tmp_path))
    second = pilot.run_spreadsheet_pilot(xlsx=path, limit=5, dry_run=True, output_root=output, client_factory=fake_client_factory(tmp_path))
    assert manifest["run_id"] != second["run_id"]
    assert second["base_run_id"] == manifest["base_run_id"]
    assert manifest["database_connected"] is True
    assert manifest["products_count"] == 5
    proposal_path = output / manifest["run_id"] / manifest["proposals"][0]
    proposal = json.loads(proposal_path.read_text(encoding="utf-8"))
    assert proposal["commercial_price"] is None
    assert proposal["price_brl_reference"]
    assert proposal["supplier_images"]
    assert proposal["supplier_images"][0]["status"] == "stored"
    assert proposal["review_required"] is True
